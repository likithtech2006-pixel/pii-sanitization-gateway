"""Build the ground-truth JSON for the three Cadence sample documents.

Method (kept simple on purpose):
  1. Pull every text unit out of each source file (PPTX shapes/tables, DOCX
     paragraphs/tables/headers, PDF pages from clean 300 dpi OCR text).
  2. Harvest structured PII with strict regexes (email, phone, employee id,
     government id, date of birth).
  3. Harvest person names from a registry that is derived from the e-mail
     addresses (initial.surname@ -> "First Surname" / "F. Surname").
  4. Text that only exists inside figures / embedded images is read from
     separate OCR folders and labelled modality=ocr_figure / embedded_image.
     Structured items there are required; names there are optional.
  5. Known non-PII look-alikes (policy ids, case numbers, ...) are written
     as hard negatives.

The output holds raw (synthetic) PII, so it is git-ignored by design.

Usage:
  python -m evaluation.tools.build_ground_truth --pptx X.pptx --docx Y.docx \
      --ocr-dir pdfocr_final --fig-ocr-dir pdfocr --docx-img-ocr dxocr \
      --pdf-name Z.pdf --out evaluation/ground_truth/local/cadence_gold.json
"""
import argparse
import collections
import json
import os
import re
import unicodedata

EMAIL_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._%+-]*@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+")
PHONE_RE = re.compile(r"(?<![\w/])(?:\+\d{1,3}[ .-]?(?:\(\d{1,4}\)|\d{1,5})(?:[ .-]?\d{2,5}){2,4}|\(\d{3}\)[ ]?555-\d{4})(?![\w/])")
SHORT_PHONE_RE = re.compile(r"(?<![\w/(])(?:M )?\+1 \(555\) \d{4}(?!\d)")
EMP_RE = re.compile(r"\bEMP-\s?\d{5}\b")
SSN_RE = re.compile(r"\b\d{3}-\d{2}-\d{4}\b")
PASSPORT_RE = re.compile(r"\b[A-Z]{2}\d{7}\b")
DOB_RE = re.compile(r"(?:DOB|date of birth|born)[:,]?\s+(\d{1,2} [A-Z][a-z]+ \d{4})", re.I)

PAN_RE = re.compile(r"\b[A-Z]{5}\d{4}[A-Z]\b")
MER_RE = re.compile(r"\bMER-[A-Z]{2}-\d{4}\b")
DIR_RE = re.compile(r"\bDIR-\d{4}\b")
SSN4_RE = re.compile(r"social security number \((\d{4})\)", re.I)
ENDING_RE = re.compile(r"(?:card|account) ending (\d{4})", re.I)
ADDR_RE = re.compile(r"\b\d{1,5} [A-Z][A-Za-z]+(?: [A-Z][A-Za-z]+)* (?:Street|Lane|Court|Crescent|Road|Avenue|Drive|Way|Boulevard)"
                     r"(?:, (?:Apartment|Apt|Suite) \w+)?, [A-Z][a-z]+(?: [A-Z][a-z]+)*, [A-Z]{2} \d{5}")
CORP_ADDR = "245 Hudson Street"

HARD_NEG_RE = re.compile(r"\b(?:POL|ISS|INC|RSK|BRC|SPK|LOS|EXC|CHL|ACC|PH|CTL|AUD|REQ)-\d{3,4}(?:-\d{2,4})?\b")


def ascii_fold(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()


def key_of(s):
    return re.sub(r"[^a-z]", "", ascii_fold(s).lower())


# ---------------------------------------------------------------- extraction
def pptx_units(path):
    from pptx import Presentation
    units = []
    prs = Presentation(path)
    for n, slide in enumerate(prs.slides, 1):
        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6:  # group
                    walk(sh.shapes)
                    continue
                if sh.has_text_frame and sh.text_frame.text.strip():
                    units.append(({"slide": n}, sh.text_frame.text, "native_text"))
                if getattr(sh, "has_table", False) and sh.has_table:
                    for r, row in enumerate(sh.table.rows, 1):
                        for cell in row.cells:
                            if cell.text.strip():
                                units.append(({"slide": n, "row": r}, cell.text, "table"))
        walk(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            units.append(({"slide": n}, slide.notes_slide.notes_text_frame.text, "native_text"))
    return units


def docx_units(path):
    import docx
    d = docx.Document(path)
    units = []
    for p in d.paragraphs:
        if p.text.strip():
            units.append(({}, p.text, "narrative"))
    for t in d.tables:
        for r, row in enumerate(t.rows, 1):
            for cell in row.cells:
                if cell.text.strip():
                    units.append(({"row": r}, cell.text, "table"))
    for s in d.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs:
                if p.text.strip():
                    units.append(({}, p.text, "narrative"))
    cp = d.core_properties
    for field in ("author", "last_modified_by"):
        v = getattr(cp, field, "")
        if v and v.strip():
            units.append(({}, v, "metadata"))
    return units


def ocr_units(folder, modality):
    units = []
    for name in sorted(os.listdir(folder)):
        m = re.match(r"p-(\d+)\.png\.txt$", name)
        if m:
            with open(os.path.join(folder, name), encoding="utf-8", errors="ignore") as fh:
                units.append(({"page": int(m.group(1))}, fh.read(), modality))
    return units


def image_ocr_units(folder):
    units = []
    for name in sorted(os.listdir(folder)):
        if name.endswith(".txt"):
            with open(os.path.join(folder, name), encoding="utf-8", errors="ignore") as fh:
                units.append(({"image": name[:-4]}, fh.read(), "embedded_image"))
    return units


# ------------------------------------------------------------------ harvest
def harvest_structured(units):
    found = []  # (type, value, location, modality)
    for loc, text, mod in units:
        flat = re.sub(r"[ \t]+", " ", text)
        for m in EMAIL_RE.finditer(flat):
            found.append(("EMAIL", m.group(0).rstrip(".,;"), loc, mod))
        for m in PHONE_RE.finditer(flat):
            digits = re.sub(r"\D", "", m.group(0))
            if 9 <= len(digits) <= 15:
                found.append(("PHONE?" if "(800)" in m.group(0) else "PHONE", m.group(0).strip(), loc, mod))
        for m in SHORT_PHONE_RE.finditer(flat):
            found.append(("PHONE", re.sub(r"^M ", "", m.group(0)), loc, mod))
        for m in EMP_RE.finditer(flat):
            found.append(("EMPLOYEE_ID", re.sub(r"\s", "", m.group(0)), loc, mod))
        for m in SSN_RE.finditer(flat):
            found.append(("GOVERNMENT_ID", m.group(0), loc, mod))
        for m in PASSPORT_RE.finditer(flat):
            found.append(("GOVERNMENT_ID", m.group(0), loc, mod))
        for m in PAN_RE.finditer(flat):
            found.append(("GOVERNMENT_ID", m.group(0), loc, mod))
        for m in MER_RE.finditer(flat):
            found.append(("EMPLOYEE_ID", m.group(0), loc, mod))
        for m in DIR_RE.finditer(flat):
            found.append(("EMPLOYEE_ID?", m.group(0), loc, mod))   # '?' = optional
        for m in SSN4_RE.finditer(flat):
            found.append(("GOVERNMENT_ID?", m.group(1), loc, mod))
        for m in ENDING_RE.finditer(flat):
            found.append(("FINANCIAL?", m.group(1), loc, mod))
        for m in ADDR_RE.finditer(flat):
            found.append(("ADDRESS?" if CORP_ADDR in m.group(0) else "ADDRESS", m.group(0), loc, mod))
        for m in DOB_RE.finditer(flat):
            found.append(("DOB", m.group(1), loc, mod))
    return found


NAME_TOKEN = r"[A-Z][A-Za-z'\u2019\u00C0-\u017F-]+"
FULL_NAME_RE = re.compile(r"\b(%s(?: [A-Z]\.)?(?: %s){1,2})\b" % (NAME_TOKEN, NAME_TOKEN))
EMP_NAME_RE = re.compile(r"(%s(?: [A-Z]\.)?(?: %s)+),? ?(?:\(|\u2014 |- |: )?EMP-" % (NAME_TOKEN, NAME_TOKEN))


STOP_WORDS = {"limited", "europe", "officer", "counsel", "framework", "engineer", "risk", "chief",
              "group", "policy", "annual", "standard", "medium", "high", "board", "committee",
              "financial", "privacy", "security", "data", "sa", "inc", "ltd", "pte", "singapore"}


def is_person_like(full):
    words = full.split(" ")
    if any(w.lower().strip(".") in STOP_WORDS for w in words):
        return False
    return not (words[0].isupper() and len(words[0]) > 1)


def split_name(full):
    parts = full.split(" ")
    mid = [p for p in parts[1:-1] if re.match(r"^[A-Z]\.$", p)]
    return parts[0], (mid[0] if mid else ""), parts[-1]


def person_registry(all_units):
    """People known from (a) e-mail local parts and (b) 'Name (EMP-nnnnn)'.

    Returns (registry, locals_) where registry is {canonical full name: (first, middle, last)}.
    """
    emails = set()
    for _, text, _ in all_units:
        for m in EMAIL_RE.finditer(text):
            emails.add(m.group(0).lower())
    locals_ = {}
    for e in emails:
        local = e.split("@")[0].lstrip("._")
        if re.match(r"^[a-z]\.[a-z]+$", local):
            locals_[local] = (local[0], local[2:])
    reg = {}
    for _, text, _ in all_units:
        flat = re.sub(r"\s+", " ", text)
        for m in EMP_NAME_RE.finditer(flat):
            full = m.group(1)
            words = full.split(" ")
            inits = [i for i, w in enumerate(words) if re.match(r"^[A-Z]\.$", w)]
            if inits and inits[-1] >= 1:
                words = words[inits[-1] - 1:]      # "First M. Last"
            else:
                words = words[-2:]                 # "First Last"
            full = " ".join(words)
            f, mid, l = split_name(full)
            if len(words) >= 2 and len(f) > 1 and is_person_like(full):
                reg[full] = (f, mid, l)
        for m in FULL_NAME_RE.finditer(flat):
            words = m.group(1).split(" ")
            for start in range(len(words) - 1):
                for end in (start + 2, start + 3):
                    cand = words[start:end]
                    if len(cand) < 2:
                        continue
                    f, mid, l = split_name(" ".join(cand))
                    fk, lk = key_of(f), key_of(l)
                    if len(fk) < 2 or not lk:
                        continue
                    for ini, sur in locals_.values():
                        if fk.startswith(ini) and (lk == sur or (len(sur) >= 5 and lk.startswith(sur))):
                            if is_person_like(" ".join(cand)):
                                reg[" ".join(cand)] = (f, mid, l)
    clean = {}
    for full, (f, mid, l) in reg.items():
        full = re.sub(r"['\u2019]s$", "", full)
        l = re.sub(r"['\u2019]s$", "", l)
        fk, lk = key_of(f), key_of(l)
        owners = [ini for ini, sur in locals_.values() if lk == sur or (len(sur) >= 5 and lk.startswith(sur))]
        if owners and not any(fk.startswith(o) for o in owners):
            continue
        clean[full] = (f, mid, l)
    return clean, locals_


def surfaces(f, mid, l):
    out = {"%s %s" % (f, l), "%s. %s" % (f[0], l), "%s %s" % (f[0], l)}
    if mid:
        out.add("%s %s %s" % (f, mid, l))
        out.add("%s. %s %s" % (f[0], mid, l))
    return out


def harvest_names(units, reg, locals_):
    # one entry per distinct surface string
    surf = {}
    for full, (f, mid, l) in reg.items():
        for sfc in surfaces(f, mid, l) | {full}:
            surf[sfc] = full
    ordered = sorted(surf, key=len, reverse=True)
    found = []
    for loc, text, mod in units:
        flat = re.sub(r"\s+", " ", text)
        covered = []
        for sfc in ordered:
            for m in re.finditer(r"(?<![A-Za-z])%s(?![A-Za-z.])|(?<![A-Za-z])%s(?![A-Za-z])" % (re.escape(sfc), re.escape(sfc)), flat):
                if any(s <= m.start() < e for s, e in covered):
                    continue
                covered.append((m.start(), m.end()))
                found.append(("PERSON", sfc, loc, mod))
    return found


def hard_negatives(units):
    seen = {}
    for loc, text, _ in units:
        for m in HARD_NEG_RE.finditer(text):
            seen.setdefault(m.group(0), loc)
    return [{"text": t, "why": "business/case/policy reference, not personal data", "location": l}
            for t, l in sorted(seen.items())]


# -------------------------------------------------------------------- build
def add_extra_people(reg, names):
    for full in names:
        words = full.split(" ")
        f, mid, l = split_name(full) if len(words) > 1 else ("", "", full)
        reg[full] = (f, mid, l)


def build_doc(key, source_name, units, names, locals_, fig_units=None, fig_mod="ocr_figure"):
    items, n = [], 0
    known_locals = {e.split('@')[0].lower() for _, t, _ in units for e in EMAIL_RE.findall(t) if e.lower().endswith(('cadencefg.com', '.example'))}
    known_locals |= {'%s.%s' % v for v in locals_.values()}

    def add(typ, value, loc, mod, required, ver="text", note=""):
        nonlocal n
        if typ.endswith("?"):
            typ, required = typ[:-1], False
        if typ == "EMAIL":
            local, _, dom = value.lower().partition("@")
            if not re.search(r"(\.example|^cadencefg\.(com|example))$", dom):
                if local in known_locals:           # OCR damaged a known Cadence address
                    value = local + "@cadencefg.com"
                else:
                    return                          # OCR garbage
        n += 1
        items.append({"id": "%s-%04d" % (key, n), "type": typ, "value": value,
                      "location": loc, "required": required, "modality": mod,
                      "verification": ver, "note": note})

    base = harvest_structured(units) + harvest_names(units, names, locals_)
    for typ, val, loc, mod in base:
        add(typ, val, loc, mod, True)
    if fig_units:
        have = collections.Counter((t, v.lower(), tuple(sorted(l.items()))) for t, v, l, _ in base)
        extra = harvest_structured(fig_units) + harvest_names(fig_units, names, locals_)
        for typ, val, loc, _ in extra:
            k = (typ, val.lower(), tuple(sorted(loc.items())))
            if have[k] > 0:           # already counted from the clean text of the same page
                have[k] -= 1
                continue
            structured = typ != "PERSON"
            add(typ, val, loc, fig_mod, structured, "text",
                "read from OCR of a figure/image" + ("" if structured else "; name OCR, optional"))
    return {"source_name": source_name, "source_type": key, "items": items,
            "hard_negatives": hard_negatives(units)}


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--docx", required=True)
    ap.add_argument("--pdf-name", required=True)
    ap.add_argument("--ocr-dir", required=True, help="clean OCR text, figure regions masked")
    ap.add_argument("--fig-ocr-dir", required=True, help="OCR of full pages (incl. figures)")
    ap.add_argument("--docx-img-ocr", required=True)
    ap.add_argument("--manual", help="optional JSON of manually verified extra items")
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)

    pptx_u = pptx_units(a.pptx)
    docx_u = docx_units(a.docx)
    pdf_u = ocr_units(a.ocr_dir, "ocr_text")
    pdf_full = ocr_units(a.fig_ocr_dir, "ocr_figure")
    img_u = image_ocr_units(a.docx_img_ocr)

    names, locals_ = person_registry(pptx_u + docx_u + pdf_u + pdf_full + img_u)
    manual = {}
    if a.manual:
        with open(a.manual, encoding="utf-8") as fh:
            manual = json.load(fh)
    add_extra_people(names, manual.get("_people", []))
    docs = {
        "pptx": build_doc("pptx", os.path.basename(a.pptx), pptx_u, names, locals_),
        "docx": build_doc("docx", os.path.basename(a.docx), docx_u, names, locals_,
                          img_u, "embedded_image"),
        "pdf": build_doc("pdf", a.pdf_name, pdf_u, names, locals_, pdf_full, "ocr_figure"),
    }
    if manual:
        for key, extra in manual.items():
            if key.startswith("_"):
                continue
            for it in extra:
                it.setdefault("id", "%s-m%03d" % (key, len(docs[key]["items"]) + 1))
                it.setdefault("verification", "visual")
                docs[key]["items"].append(it)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with open(a.out, "w", encoding="utf-8") as fh:
        json.dump({"schema_version": 1, "documents": docs}, fh, indent=1, ensure_ascii=False)
    for k, d in docs.items():
        req = sum(1 for i in d["items"] if i["required"])
        print("%-5s items=%4d required=%4d hard_negatives=%d" % (k, len(d["items"]), req, len(d["hard_negatives"])))
    print("people in registry:", len(names))


if __name__ == "__main__":
    main()
