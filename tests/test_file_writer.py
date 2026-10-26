"""Run from repo root:  pytest tests/test_file_writer.py -v
Format-preserving redacted copies (fake data only)."""
import io
import zipfile

import pytest
from PIL import Image, ImageDraw, ImageFont

import detection
from extraction import extract, ocr
from redaction import sanitize_and_verify
from redaction.file_writer import redacted_spans, write_redacted_file

try:
    import pymupdf
except ImportError:
    import fitz as pymupdf

FAKE = ["Jane Smith", "jane.smith@example.com", "+1 (555) 0100", "John Doe", "900-42-7781"]


def run(path):
    doc = extract(str(path))
    det = detection.PIIDetectorImpl()
    gate = sanitize_and_verify(doc, det.detect(doc), det)
    return doc, gate


def all_xml_text(path):
    with zipfile.ZipFile(path) as z:
        return "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith((".xml", ".rels")))


def test_redacted_spans_alignment():
    spans = redacted_spans("Call Jane Smith at 555-0100.", "Call [PERSON_1] at [PHONE_1].")
    assert [(s, e, t) for s, e, t in spans] == [(5, 15, "[PERSON_1]"), (19, 27, "[PHONE_1]")]


def test_docx_copy(tmp_path):
    from docx import Document
    d = Document()
    p = d.add_paragraph("Owner: ")
    p.add_run("Jane").bold = True
    p.add_run(" Smith, jane.smith@example.com")          # a name split over two runs
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Phone"
    t.cell(1, 0).text, t.cell(1, 1).text = "John Doe", "+1 (555) 0100"
    d.sections[0].header.paragraphs[0].text = "Prepared for Jane Smith"
    d.core_properties.author = "John Doe"
    src = tmp_path / "in.docx"; d.save(src)

    doc, gate = run(src)
    out = tmp_path / "out.docx"
    outcome = write_redacted_file(str(src), doc, gate, str(out))
    assert outcome.ok, outcome.reason
    xml = all_xml_text(out)
    for raw in FAKE[:4]:
        assert raw not in xml
    from docx import Document as Open
    text = "\n".join(p.text for p in Open(out).paragraphs)
    assert "[PERSON_" in text and "[EMAIL_" in text


def test_pptx_copy(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Leadership"
    tf = s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame
    tf.text = "John Doe"
    tf.add_paragraph().text = "M +1 (555) 0100"
    tf.add_paragraph().text = "jane.smith@example.com"
    s.notes_slide.notes_text_frame.text = "Ask Jane Smith"
    src = tmp_path / "in.pptx"; prs.save(src)

    doc, gate = run(src)
    out = tmp_path / "out.pptx"
    outcome = write_redacted_file(str(src), doc, gate, str(out))
    assert outcome.ok, outcome.reason
    xml = all_xml_text(out)
    template_editor = Presentation().core_properties.last_modified_by   # a name in python-pptx's template
    for raw in ["John Doe", "+1 (555) 0100", "jane.smith@example.com", "Jane Smith", template_editor]:
        assert raw not in xml                             # incl. notes and file properties
    Presentation(out)                                     # still a valid file


def test_text_pdf_copy(tmp_path):
    d = pymupdf.open()
    page = d.new_page()
    page.insert_text((72, 100), "Owner: Jane Smith, jane.smith@example.com, SSN 900-42-7781", fontsize=11)
    page.insert_text((72, 130), "The committee meets quarterly.", fontsize=11)
    src = tmp_path / "in.pdf"; d.save(src)

    doc, gate = run(src)
    out = tmp_path / "out.pdf"
    outcome = write_redacted_file(str(src), doc, gate, str(out))
    assert outcome.ok, outcome.reason
    text = pymupdf.open(out)[0].get_text()
    for raw in ["Jane Smith", "jane.smith@example.com", "900-42-7781"]:
        assert raw not in text
    assert "The committee meets quarterly." in text        # everything else is kept


@pytest.mark.skipif(not ocr.is_available(), reason="Tesseract not installed")
def test_scanned_pdf_copy_is_verified_by_ocr(tmp_path):
    img = Image.new("RGB", (1700, 500), "white")
    font = ImageFont.load_default(size=50)
    ImageDraw.Draw(img).text((60, 100), "Owner Jane Smith", fill="black", font=font)
    ImageDraw.Draw(img).text((60, 250), "Email jane.smith@example.com", fill="black", font=font)
    buf = io.BytesIO(); img.save(buf, "PNG")
    d = pymupdf.open(); pg = d.new_page(width=612, height=180)
    pg.insert_image(pg.rect, stream=buf.getvalue())
    src = tmp_path / "scan.pdf"; d.save(src)

    doc, gate = run(src)
    assert gate.is_safe
    out = tmp_path / "scan.redacted.pdf"
    outcome = write_redacted_file(str(src), doc, gate, str(out))
    assert outcome.ok, outcome.reason
    assert outcome.stats["ocr_rounds"] >= 1
    check = extract(str(out))
    assert "smith" not in check.full_text().lower().replace("[person", "")


def test_blocked_document_is_not_written(tmp_path):
    from docx import Document
    d = Document(); d.add_paragraph("Reach jane.smith@example.com")
    src = tmp_path / "leak.docx"; d.save(src)
    doc = extract(str(src))
    gate = sanitize_and_verify(doc, [], max_passes=1)
    outcome = write_redacted_file(str(src), doc, gate, str(tmp_path / "x.docx"))
    assert not outcome.ok and not (tmp_path / "x.docx").exists()
