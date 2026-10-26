"""Run from repo root:  pytest tests/test_extraction.py -v
Fixtures are generated on the fly (fake PII only), so no real files are needed.
Real-file tests at the bottom run only if you put files in samples/ locally."""
import dataclasses
import io
import json
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFont

from extraction import extract, ocr

try:
    import pymupdf
except ImportError:
    import fitz as pymupdf

needs_ocr = pytest.mark.skipif(not ocr.is_available(), reason="Tesseract not installed")


# ---------- fixtures (fake data only) ----------
@pytest.fixture
def docx_file(tmp_path):
    from docx import Document
    d = Document()
    d.add_paragraph("Customer: Jane Smith, email jane.smith@example.com")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Phone"
    t.cell(1, 0).text, t.cell(1, 1).text = "John Doe", "+1 555-010-9999"
    d.add_paragraph("")                                  # empty paragraph must be skipped
    d.add_paragraph("Closing paragraph")
    d.sections[0].footer.paragraphs[0].text = "Footer: acct 123456789"
    p = tmp_path / "t.docx"; d.save(p); return p


@pytest.fixture
def pptx_file(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Quarterly Review"
    s1.placeholders[1].text = "Contact: jane.smith@example.com"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Second slide"
    s2.notes_slide.notes_text_frame.text = "Speaker note: call 555-0100"
    p = tmp_path / "t.pptx"; prs.save(p); return p


@pytest.fixture
def text_pdf(tmp_path):
    doc = pymupdf.open()
    for n in (1, 2):
        pg = doc.new_page()
        pg.insert_text((72, 100), f"Page {n}: contact jane.smith@example.com for details", fontsize=12)
    p = tmp_path / "t.pdf"; doc.save(p); return p


@pytest.fixture
def scanned_pdf(tmp_path):
    img = Image.new("RGB", (1700, 600), "white")
    font = ImageFont.load_default(size=50)
    ImageDraw.Draw(img).text((60, 100), "Customer John Smith", fill="black", font=font)
    ImageDraw.Draw(img).text((60, 250), "Email john.smith@example.com", fill="black", font=font)
    buf = io.BytesIO(); img.save(buf, "PNG")
    doc = pymupdf.open(); pg = doc.new_page(width=612, height=216)
    pg.insert_image(pg.rect, stream=buf.getvalue())      # image only, no text layer
    p = tmp_path / "scan.pdf"; doc.save(p); return p


def loc(b):
    return json.loads(b.metadata["location"])


def body(doc):
    """Content blocks without the document-property blocks appended at the end."""
    return [b for b in doc.content if loc(b)["type"] != "metadata"]


def common_checks(doc):
    assert doc.content, "no blocks extracted"
    assert [int(b.metadata["order"]) for b in doc.content] == list(range(len(doc.content)))
    ids = [b.metadata["block_id"] for b in doc.content]
    assert len(ids) == len(set(ids)), "duplicate block_id"
    assert all(b.source for b in doc.content), "block without source"
    assert all(b.text.strip() for b in doc.content), "empty block"
    assert doc.metadata["status"] == "ok"


# ---------- tests ----------
def test_docx(docx_file):
    doc = extract(docx_file)
    common_checks(doc)
    assert doc.source_type == "docx"
    types = {loc(b)["type"] for b in doc.content}
    assert {"paragraph", "table", "footer"} <= types
    assert any("John Doe" in b.text for b in doc.content)          # table cell captured
    assert any(b.block_type == "table_cell" for b in doc.content)
    assert loc(doc.content[0]) == {"type": "paragraph", "index": 0}


def test_pptx(pptx_file):
    doc = extract(pptx_file)
    common_checks(doc)
    assert doc.source_type == "pptx"
    assert {loc(b)["slide"] for b in body(doc)} == {1, 2}
    assert any(b.block_type == "notes" for b in doc.content)
    assert all(b.source.startswith("slide") for b in body(doc))


def test_text_pdf(text_pdf):
    doc = extract(text_pdf)
    common_checks(doc)
    assert doc.source_type == "pdf"
    assert doc.metadata["ocr_used"] == "false"
    assert {loc(b)["page"] for b in body(doc)} == {1, 2}
    assert all(json.loads(b.metadata["bbox"]) for b in body(doc))
    assert doc.content[0].source == "page 1"


@needs_ocr
def test_scanned_pdf_uses_ocr(scanned_pdf):
    doc = extract(scanned_pdf)
    common_checks(doc)
    assert doc.metadata["ocr_used"] == "true"
    assert all(b.block_type == "ocr_text" for b in body(doc))
    assert "smith" in doc.full_text().lower()
    assert all(loc(b)["page"] == 1 for b in body(doc))


def test_bytes_input(docx_file):
    doc = extract(docx_file.read_bytes(), filename="upload.docx")   # how Streamlit will call it
    assert doc.content and doc.file_name == "upload.docx"


def test_unsupported_type_does_not_crash(tmp_path):
    p = tmp_path / "x.txt"; p.write_text("hi")
    doc = extract(p)
    assert doc.content == [] and doc.metadata["status"] == "error"
    assert "unsupported" in doc.metadata["warnings"]


def test_corrupt_file_does_not_crash(tmp_path):
    p = tmp_path / "bad.pdf"; p.write_bytes(b"this is not a pdf")
    doc = extract(p)
    assert doc.content == [] and doc.metadata["status"] == "error"


def test_document_properties_are_extracted_last(tmp_path):
    from docx import Document
    d = Document()
    d.add_paragraph("Body text")
    d.core_properties.author = "Jane Smith"
    d.core_properties.last_modified_by = "John Doe"
    p = tmp_path / "props.docx"; d.save(p)
    doc = extract(p)
    props = [b for b in doc.content if loc(b)["type"] == "metadata"]
    assert {b.text for b in props} >= {"Jane Smith", "John Doe"}
    assert doc.content[-len(props):] == props                      # appended after the body
    assert all(b.block_type == "text" and b.source.startswith("metadata | field") for b in props)
    assert "python-docx" not in {b.text for b in props}            # tool names are not extracted


def test_pdf_author_is_extracted(text_pdf, tmp_path):
    doc = pymupdf.open(text_pdf)
    doc.set_metadata({"author": "Jane Smith", "producer": "SomeTool 1.0"})
    out = tmp_path / "author.pdf"; doc.save(out)
    props = [b.text for b in extract(out).content if loc(b)["type"] == "metadata"]
    assert props == ["Jane Smith"]                                  # producer is software, not a person


def test_docx_text_boxes_are_extracted(tmp_path):
    """Text boxes live inside drawings; python-docx's paragraph text does not include them."""
    from docx import Document
    from docx.oxml import parse_xml
    d = Document()
    para = d.add_paragraph("Before the box")
    box = parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        '<w:drawing><wp:anchor><a:graphic><a:graphicData><wps:wsp><wps:txbx><w:txbxContent>'
        '<w:p><w:r><w:t>Escalate to jane.smith@example.com</w:t></w:r></w:p>'
        '</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>')
    para._p.append(box)
    p = tmp_path / "box.docx"; d.save(p)
    doc = extract(p)
    boxes = [b for b in doc.content if loc(b)["type"] == "textbox"]
    assert [b.text for b in boxes] == ["Escalate to jane.smith@example.com"]


@needs_ocr
def test_docx_screenshot_is_ocrd(tmp_path):
    from docx import Document
    from docx.shared import Inches
    img = Image.new("RGB", (1400, 300), "white")
    ImageDraw.Draw(img).text((40, 100), "Reviewer Jane Smith", fill="black", font=ImageFont.load_default(size=60))
    buf = io.BytesIO(); img.save(buf, "PNG"); buf.seek(0)
    d = Document()
    d.add_paragraph("Screenshot below")
    d.add_picture(buf, width=Inches(5))
    p = tmp_path / "shot.docx"; d.save(p)
    doc = extract(p)
    pics = [b for b in doc.content if loc(b).get("image_ocr")]
    assert pics and "smith" in " ".join(b.text for b in pics).lower()
    assert all(b.block_type == "ocr_text" for b in pics)


def test_json_serialisable(docx_file):
    json.dumps(dataclasses.asdict(extract(docx_file)))


# ---------- optional: real Cadence files kept LOCALLY in samples/ (never committed) ----------
@pytest.mark.parametrize("pattern", ["*.pdf", "*.pptx", "*.docx"])
def test_real_files_if_present(pattern):
    files = list(Path("samples").glob(pattern))
    if not files:
        pytest.skip(f"no {pattern} in samples/")
    for f in files:
        doc = extract(f)
        common_checks(doc)
        print(f.name, len(doc.content), "blocks, ocr:", doc.metadata["ocr_used"])