"""PPTX -> blocks. One block per non-empty paragraph, with slide number
and shape id so Sanjay can find it again. Handles tables, groups, notes
and (optionally) OCR of pictures."""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from extraction import images, ocr
from extraction.internal import Block
from extraction.metadata import add_metadata_blocks, office_properties


def extract_pptx(path, doc_id, meta):
    prs = Presentation(path)
    meta.parser = "python-pptx"
    blocks, state = [], {"order": 0}

    def add(text, loc, md):
        if not text or not text.strip():
            return
        blocks.append(Block(f"{doc_id}-b{state['order']}", text.strip(), loc,
                            state["order"], md))
        state["order"] += 1

    def walk(shapes, slide_no):
        for shape in shapes:
            sid, sname = shape.shape_id, shape.name
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, slide_no)               # recurse into groups
            elif getattr(shape, "has_table", False) and shape.has_table:
                for r, row in enumerate(shape.table.rows):
                    for c, cell in enumerate(row.cells):
                        add(cell.text, {"type": "slide", "slide": slide_no,
                                        "shape": sname, "row": r, "col": c},
                            {"shape_id": sid, "kind": "table", "row": r, "col": c})
            elif shape.has_text_frame:
                for p_i, para in enumerate(shape.text_frame.paragraphs):
                    text = "".join(run.text for run in para.runs)
                    add(text, {"type": "slide", "slide": slide_no,
                               "shape": sname, "paragraph": p_i},
                        {"shape_id": sid, "kind": "text", "paragraph": p_i})
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                _ocr_picture(shape, slide_no, add, meta)

    for slide_no, slide in enumerate(prs.slides, 1):
        walk(slide.shapes, slide_no)
        if slide.has_notes_slide:                          # speaker notes
            for p_i, para in enumerate(slide.notes_slide.notes_text_frame.paragraphs):
                add(para.text, {"type": "slide_notes", "slide": slide_no, "paragraph": p_i},
                    {"kind": "notes", "paragraph": p_i})
    add_metadata_blocks(add, office_properties(prs.core_properties))
    return blocks


def _ocr_picture(shape, slide_no, add, meta):
    """Pictures can contain screenshots with PII. OCR them if possible."""
    try:
        img = images.load_image(shape.image.blob, shape.image.filename or "picture.png", meta)
        if img is None or min(img.size) < images.MIN_SIDE:
            return                                         # unreadable (warned) or just an icon/logo
        if not ocr.is_available():
            meta.warnings.append(f"slide {slide_no}: picture skipped, Tesseract is not installed")
            return
        for para in ocr.ocr_image(img, dpi=72 * 4):
            add(para["text"], {"type": "slide", "slide": slide_no,
                               "shape": shape.name, "image_ocr": True},
                {"shape_id": shape.shape_id, "kind": "picture_ocr", "ocr": True,
                 "ocr_conf": para["conf"]})
            meta.ocr_used = True
    except Exception as e:                                 # never crash on one picture
        meta.warnings.append(f"slide {slide_no}: picture OCR failed ({type(e).__name__})")
