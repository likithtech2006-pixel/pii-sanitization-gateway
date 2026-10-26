"""DOCX -> blocks. Body paragraphs and tables in true document order,
text boxes, headers/footers, footnotes/endnotes/comments (PII often hides
there), OCR of embedded pictures (screenshots), and document properties."""
from concurrent.futures import ThreadPoolExecutor

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from extraction import images, ocr
from extraction.internal import Block
from extraction.metadata import add_metadata_blocks, office_properties
from security import xmlsafe

_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_VML_IMAGE = "{urn:schemas-microsoft-com:vml}imagedata"
_VML_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_NOTE_PARTS = (("footnotes", "/word/footnotes.xml"), ("endnotes", "/word/endnotes.xml"),
               ("comments", "/word/comments.xml"))


def _textbox_paragraphs(element):
    """Paragraph texts inside text boxes of an element (each box appears once even though
    Word stores a modern and a legacy VML copy)."""
    seen, out = set(), []
    for box in element.iter(_W + "txbxContent"):
        for p in box.iter(_W + "p"):
            text = "".join(t.text or "" for t in p.iter(_W + "t"))
            if text.strip() and text not in seen:
                seen.add(text)
                out.append(text)
    return out


def _image_rel_ids(element):
    """Relationship ids of every picture inside an XML element, in order."""
    ids = []
    for blip in element.iter(_BLIP):
        if blip.get(_EMBED):
            ids.append(blip.get(_EMBED))
    for vml in element.iter(_VML_IMAGE):
        if vml.get(_VML_ID):
            ids.append(vml.get(_VML_ID))
    return ids


def extract_docx(path, doc_id, meta):
    d = Document(path)
    meta.parser = "python-docx"
    blocks, state = [], {"order": 0}

    def add(text, loc, md=None):
        if not text or not text.strip():
            return
        blocks.append(Block(f"{doc_id}-b{state['order']}", text.strip(), loc,
                            state["order"], md or {}))
        state["order"] += 1

    # Pictures are OCR'd in parallel up front, then emitted where they sit in the body.
    pictures = _ocr_pictures(d, meta)
    emitted = set()

    def add_pictures(element, loc):
        for rel_id in _image_rel_ids(element):
            if rel_id in emitted or rel_id not in pictures:
                continue                                   # same picture used twice -> once
            emitted.add(rel_id)
            name, paras = pictures[rel_id]
            for para in paras:
                add(para["text"], dict(loc, image=name, image_ocr=True),
                    {"kind": "picture_ocr", "ocr": True, "ocr_conf": para["conf"], "image": name})
                meta.ocr_used = True

    p_i = t_i = 0
    for child in d.element.body.iterchildren():
        if child.tag.endswith("}p"):                       # paragraph
            add(Paragraph(child, d).text,
                {"type": "paragraph", "index": p_i}, {"body_paragraph_index": p_i})
            for k, box_text in enumerate(_textbox_paragraphs(child)):
                add(box_text, {"type": "textbox", "paragraph": p_i, "index": k}, {"kind": "textbox"})
            add_pictures(child, {"type": "paragraph", "index": p_i})
            p_i += 1
        elif child.tag.endswith("}tbl"):                   # table
            table, seen = Table(child, d), set()
            for r, row in enumerate(table.rows):
                for c, cell in enumerate(row.cells):
                    if cell._tc in seen:                   # merged cell repeats
                        continue
                    seen.add(cell._tc)
                    add(cell.text, {"type": "table", "table": t_i, "row": r, "col": c},
                        {"table_index": t_i, "row": r, "col": c})
            add_pictures(child, {"type": "table", "table": t_i})
            t_i += 1

    for s_i, section in enumerate(d.sections):             # headers / footers
        for kind, part in (("header", section.header), ("footer", section.footer)):
            for j, para in enumerate(part.paragraphs):
                add(para.text, {"type": kind, "section": s_i, "index": j},
                    {"section": s_i, "part": kind, "paragraph_index": j})
            for k, table in enumerate(part.tables):
                for r, row in enumerate(table.rows):
                    for c, cell in enumerate(row.cells):
                        add(cell.text, {"type": kind + "_table", "section": s_i,
                                        "table": k, "row": r, "col": c})

            for k, box_text in enumerate(_textbox_paragraphs(part._element)):
                add(box_text, {"type": kind + "_textbox", "section": s_i, "index": k}, {"kind": "textbox"})

    # Footnotes, endnotes and comments live in their own parts.
    parts = {str(p.partname): p for p in d.part.package.iter_parts()}
    for kind, name in _NOTE_PARTS:
        part = parts.get(name)
        if part is None:
            continue
        try:
            root = xmlsafe.fromstring(part.blob)
        except Exception:
            meta.warnings.append(f"{kind} could not be read")
            continue
        for k, p in enumerate(root.iter(_W + "p")):
            add("".join(t.text or "" for t in p.iter(_W + "t")), {"type": kind, "index": k}, {"kind": kind})

    # Tracked changes: python-docx's paragraph.text skips <w:ins> and <w:del>, yet both stay stored in the
    # file (deleted text included). Extract them so PII inside tracked changes is detected and redacted.
    roots = [d.element.body] + [part._element for sec in d.sections for part in (sec.header, sec.footer)]
    k, seen = 0, set()
    for root in roots:
        for el in root.iter(_W + "ins", _W + "del"):
            deleted = el.tag == _W + "del"
            text = "".join(t.text or "" for t in el.iter(_W + ("delText" if deleted else "t")))
            if text.strip() and id(el) not in seen:
                seen.add(id(el))
                add(text, {"type": "tracked_deletion" if deleted else "tracked_insertion", "index": k},
                    {"kind": "tracked_change"})
                k += 1

    add_metadata_blocks(add, office_properties(d.core_properties))
    return blocks


def _ocr_pictures(d, meta):
    """rel_id -> (image name, OCR paragraphs) for every picture in the body."""
    if not ocr.is_available():
        if _image_rel_ids(d.element.body):
            meta.warnings.append("pictures skipped: Tesseract is not installed")
        return {}
    rels = d.part.related_parts
    todo = {}
    for rel_id in dict.fromkeys(_image_rel_ids(d.element.body)):
        part = rels.get(rel_id)
        if part is None or not hasattr(part, "blob"):
            continue
        name = part.partname.split("/")[-1]
        img = images.load_image(part.blob, name, meta)
        if img is not None and min(img.size) >= images.MIN_SIDE:
            todo[rel_id] = (name, img)
    out = {}
    with ThreadPoolExecutor(max_workers=ocr.workers()) as pool:
        futures = {rid: pool.submit(ocr.ocr_image, img, 72 * 4) for rid, (_, img) in todo.items()}
        for rid, future in futures.items():
            try:
                out[rid] = (todo[rid][0], future.result())
            except Exception as e:
                meta.warnings.append(f"picture {todo[rid][0]}: OCR failed ({type(e).__name__})")
    return out
